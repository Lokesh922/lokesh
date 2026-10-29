
import streamlit as st
import pandas as pd
import plotly.graph_objects as go
import re
import io
from collections import Counter
from PyPDF2 import PdfReader

# --------------------------------------------------
# PAGE CONFIGURATION
# --------------------------------------------------

st.set_page_config(
    page_title="ResumeAI | Resume Analyzer",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="expanded"
)

# --------------------------------------------------
# CUSTOM CSS
# --------------------------------------------------

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');

.stApp {
    background: #0b0d12;
    color: #eef2ff;
    font-family: 'Inter', sans-serif;
}

header[data-testid="stHeader"] {
    background: #0b0d12;
}

.block-container {
    max-width: 1450px;
    padding-top: 2rem;
    padding-bottom: 3rem;
}

[data-testid="stSidebar"] {
    background: #11141d;
    border-right: 1px solid #252a38;
}

h1, h2, h3 {
    color: #f8fafc !important;
    font-weight: 700 !important;
}

p, label, li {
    color: #b8c0d4;
}

.hero {
    padding: 24px 28px;
    border: 1px solid #292f40;
    border-radius: 18px;
    background: linear-gradient(120deg, #171c2b, #11141d);
    margin-bottom: 25px;
}

.hero-tag {
    color: #a5b4fc;
    font-size: 12px;
    font-weight: 700;
    letter-spacing: 2px;
    text-transform: uppercase;
}

.hero-title {
    font-size: 34px;
    font-weight: 800;
    color: #f8fafc;
    margin: 10px 0;
}

.hero-subtitle {
    color: #a5adc2;
    font-size: 14px;
}

.panel {
    background: #131722;
    border: 1px solid #282e3e;
    border-radius: 15px;
    padding: 20px;
    margin-bottom: 15px;
    min-height: 100px;
}

.panel-title {
    color: #f3f5fb;
    font-size: 16px;
    font-weight: 700;
    margin-bottom: 8px;
}

.panel-description {
    color: #8e98b0;
    font-size: 12px;
    margin-bottom: 15px;
}

.metric-card {
    background: #131722;
    border: 1px solid #292f40;
    border-radius: 14px;
    padding: 20px;
    min-height: 110px;
}

.metric-label {
    font-size: 12px;
    color: #939db5;
    margin-bottom: 10px;
}

.metric-value {
    color: #f8fafc;
    font-size: 29px;
    font-weight: 800;
}

.section-heading {
    color: #f8fafc;
    font-size: 23px;
    font-weight: 750;
    margin-top: 25px;
    margin-bottom: 5px;
}

.section-subtitle {
    color: #8e98b0;
    font-size: 13px;
    margin-bottom: 20px;
}

div.stButton > button {
    background: linear-gradient(90deg, #7567ff, #6255e7);
    color: white;
    border: none;
    border-radius: 9px;
    min-height: 46px;
    font-weight: 700;
    transition: 0.2s ease;
}

div.stButton > button:hover {
    background: #8b80ff;
    color: white;
    border: none;
    transform: translateY(-1px);
}

[data-testid="stFileUploader"] {
    background: #0e111a;
    border: 1px dashed #3b4560;
    border-radius: 12px;
    padding: 12px;
}

[data-testid="stFileUploader"] section {
    background: transparent;
}

textarea, input {
    background: #0e111a !important;
    color: #f8fafc !important;
}

.stTextArea textarea {
    border: 1px solid #30374a;
    border-radius: 10px;
}

[data-baseweb="tab-list"] {
    gap: 8px;
    border-bottom: 1px solid #282e3e;
}

[data-baseweb="tab"] {
    color: #aab3c7;
}

[data-baseweb="tab-highlight"] {
    background: #8175ff;
}

div[data-testid="stAlert"] {
    border-radius: 10px;
}

hr {
    border-color: #292f40;
}

.footer {
    color: #737e95;
    font-size: 12px;
    text-align: center;
    padding: 25px 0 5px 0;
}
</style>
""", unsafe_allow_html=True)


# --------------------------------------------------
# SKILLS DATABASE
# --------------------------------------------------

SKILLS_DB = {
    "Programming": [
        "python", "java", "javascript", "typescript", "c++", "c",
        "sql", "r", "go", "php", "kotlin", "swift"
    ],
    "Web Development": [
        "html", "css", "react", "angular", "vue", "node.js",
        "express", "fastapi", "django", "flask", "rest api",
        "bootstrap", "next.js"
    ],
    "Data Science & AI": [
        "machine learning", "deep learning", "artificial intelligence",
        "data science", "pandas", "numpy", "scikit-learn",
        "tensorflow", "pytorch", "nlp", "computer vision",
        "generative ai", "llm", "natural language processing"
    ],
    "Databases": [
        "mysql", "postgresql", "mongodb", "oracle", "sqlite",
        "redis", "database management"
    ],
    "Cloud & DevOps": [
        "aws", "azure", "gcp", "docker", "kubernetes",
        "git", "github", "jenkins", "linux", "ci/cd"
    ],
    "Professional Skills": [
        "communication", "leadership", "problem solving",
        "teamwork", "agile", "scrum", "project management",
        "critical thinking"
    ]
}

ALL_SKILLS = sorted({
    skill
    for skills in SKILLS_DB.values()
    for skill in skills
})


# --------------------------------------------------
# TEXT EXTRACTION
# --------------------------------------------------

def extract_resume_text(uploaded_file):
    filename = uploaded_file.name.lower()

    try:
        if filename.endswith(".pdf"):
            reader = PdfReader(uploaded_file)

            if reader.is_encrypted:
                try:
                    reader.decrypt("")
                except Exception:
                    return None, "This PDF is password-protected."

            text = ""

            for page in reader.pages:
                text += (page.extract_text() or "") + "\n"

            if not text.strip():
                return None, (
                    "No readable text was found. "
                    "This may be a scanned PDF; OCR is required."
                )

            return text, None

        elif filename.endswith(".docx"):
            from docx import Document

            document = Document(uploaded_file)

            text = "\n".join(
                paragraph.text for paragraph in document.paragraphs
            )

            for table in document.tables:
                for row in table.rows:
                    text += "\n" + " ".join(
                        cell.text for cell in row.cells
                    )

            if not text.strip():
                return None, "No readable text was found in the DOCX."

            return text, None

        return None, "Please upload a PDF or DOCX file."

    except Exception as error:
        return None, f"Could not read the file: {error}"


# --------------------------------------------------
# KEYWORD MATCHING
# --------------------------------------------------

def find_skills(text):
    text = text.lower()
    found = []

    for skill in ALL_SKILLS:
        pattern = r"(?<![a-z0-9])" + re.escape(skill.lower()) + r"(?![a-z0-9])"

        if re.search(pattern, text):
            found.append(skill)

    return sorted(set(found))


def calculate_analysis(resume_text, job_description):
    resume_skills = find_skills(resume_text)
    job_skills = find_skills(job_description)

    matched = sorted(set(resume_skills) & set(job_skills))
    missing = sorted(set(job_skills) - set(resume_skills))

    if job_skills:
        score = round(len(matched) / len(job_skills) * 100)
    else:
        score = 0

    # Resume keyword coverage is not an official ATS score.
    resume_words = set(re.findall(r"\b[a-zA-Z0-9+#.]+\b", resume_text.lower()))
    job_words = set(re.findall(r"\b[a-zA-Z0-9+#.]+\b", job_description.lower()))

    common_words = {
        "the", "and", "for", "with", "you", "your", "our", "are",
        "will", "this", "that", "from", "have", "has", "who", "what",
        "can", "all", "their", "about", "into", "work", "using",
        "years", "year", "role", "job", "ability", "required",
        "preferred", "skills", "experience", "knowledge"
    }

    meaningful_job_words = job_words - common_words
    matched_words = resume_words & meaningful_job_words

    word_coverage = (
        round(len(matched_words) / len(meaningful_job_words) * 100)
        if meaningful_job_words else 0
    )

    strengths = []
    warnings = []
    fixes = []

    if matched:
        strengths.append({
            "title": "Relevant skills identified",
            "description": "Your resume includes " + ", ".join(matched[:8]) + "."
        })
    else:
        warnings.append({
            "title": "No matching listed skills",
            "description": "Your resume may need more relevant skills for this job."
        })

    if missing:
        fixes.append({
            "title": "Address missing job skills",
            "description": (
                "Consider adding evidence of these skills if you genuinely have them: "
                + ", ".join(missing[:10]) + "."
            )
        })

    if len(resume_text.split()) < 150:
        warnings.append({
            "title": "Resume content may be limited",
            "description": (
                "The extracted resume is short. Check whether all pages and "
                "sections were read correctly."
            )
        })
    else:
        strengths.append({
            "title": "Resume text extracted",
            "description": f"Successfully extracted approximately {len(resume_text.split())} words."
        })

    if not re.search(r"\b(experience|project|projects|internship|education)\b",
                     resume_text, re.IGNORECASE):
        warnings.append({
            "title": "Check resume sections",
            "description": "Experience, projects, education, or other section headings may be missing or formatted unusually."
        })

    if len(resume_text.split()) > 1000:
        strengths.append({
            "title": "Detailed resume content",
            "description": "Your resume contains substantial text. Review it for relevance and clarity."
        })

    return {
        "score": score,
        "matched": matched,
        "missing": missing,
        "resume_skills": resume_skills,
        "job_skills": job_skills,
        "word_coverage": word_coverage,
        "resume_words": len(resume_text.split()),
        "job_words": len(job_description.split()),
        "strengths": strengths,
        "warnings": warnings,
        "fixes": fixes,
        "resume_text": resume_text
    }


# --------------------------------------------------
# CHARTS
# --------------------------------------------------

def create_gauge(score):
    if score >= 75:
        gauge_color = "#35d49a"
    elif score >= 50:
        gauge_color = "#f4bb58"
    else:
        gauge_color = "#ff6b79"

    fig = go.Figure(go.Indicator(
        mode="gauge+number",
        value=score,
        number={
            "suffix": "%",
            "font": {"size": 42, "color": "#f8fafc"}
        },
        title={
            "text": "ATS-style keyword match",
            "font": {"size": 14, "color": "#aab3c7"}
        },
        gauge={
            "axis": {
                "range": [0, 100],
                "tickwidth": 0,
                "tickcolor": "#77819a"
            },
            "bar": {"color": gauge_color, "thickness": 0.25},
            "bgcolor": "#131722",
            "borderwidth": 0,
            "steps": [
                {"range": [0, 50], "color": "#29202b"},
                {"range": [50, 75], "color": "#302c23"},
                {"range": [75, 100], "color": "#1d302d"}
            ],
        }
    ))

    fig.update_layout(
        height=240,
        margin=dict(l=15, r=15, t=45, b=10),
        paper_bgcolor="#131722",
        plot_bgcolor="#131722",
        font_color="#f8fafc"
    )

    return fig


def create_skill_chart(matched, missing):
    data = pd.DataFrame({
        "Category": ["Matched", "Missing"],
        "Count": [len(matched), len(missing)]
    })

    fig = go.Figure(go.Bar(
        x=data["Category"],
        y=data["Count"],
        marker_color=["#35d49a", "#ff6b79"],
        text=data["Count"],
        textposition="outside"
    ))

    fig.update_layout(
        height=270,
        paper_bgcolor="#131722",
        plot_bgcolor="#131722",
        font_color="#b8c0d4",
        margin=dict(l=15, r=15, t=25, b=20),
        yaxis=dict(
            title="Number of skills",
            gridcolor="#292f40",
            rangemode="tozero"
        ),
        xaxis=dict(title="")
    )

    return fig


# --------------------------------------------------
# HEADER AND SIDEBAR
# --------------------------------------------------

with st.sidebar:
    st.markdown("## 📄 ResumeAI")
    st.caption("Resume intelligence platform")
    st.divider()

    st.markdown("### Navigation")
    st.markdown("🏠  Dashboard")
    st.markdown("📊  Resume analysis")
    st.markdown("🎯  Keyword matching")
    st.markdown("📑  Resume preview")

    st.divider()

    st.markdown("### About")
    st.caption(
        "Analyze your resume against a job description, "
        "identify keyword gaps, and find areas to improve."
    )

    st.caption("Version 1.0 · Local analysis")


st.markdown("""
<div class="hero">
    <div class="hero-tag">AI-powered career toolkit</div>
    <div class="hero-title">ResumeAI Analyzer</div>
    <div class="hero-subtitle">
        Analyze your resume, discover skill gaps, and improve alignment
        with your target job description.
    </div>
</div>
""", unsafe_allow_html=True)


# --------------------------------------------------
# INPUT PANELS
# --------------------------------------------------

upload_col, jd_col = st.columns([1, 1], gap="large")

with upload_col:
    st.markdown("""
    <div class="panel">
        <div class="panel-title">📄 Upload your resume</div>
        <div class="panel-description">
            Upload your existing resume in PDF or DOCX format.
        </div>
    </div>
    """, unsafe_allow_html=True)

    uploaded_file = st.file_uploader(
        "Choose resume",
        type=["pdf", "docx"],
        label_visibility="collapsed"
    )

    if uploaded_file:
        st.success(f"Selected: {uploaded_file.name}")
        st.caption(f"File size: {uploaded_file.size / 1024:.1f} KB")

with jd_col:
    st.markdown("""
    <div class="panel">
        <div class="panel-title">🎯 Job description</div>
        <div class="panel-description">
            Paste the job requirements you want to compare against.
        </div>
    </div>
    """, unsafe_allow_html=True)

    job_description = st.text_area(
        "Job description",
        height=150,
        placeholder=(
            "Paste the job description here...\n\n"
            "Example: Python, SQL, Machine Learning, Git, "
            "communication skills..."
        ),
        label_visibility="collapsed"
    )


analyze_col, clear_col = st.columns([4, 1])

with analyze_col:
    analyze_clicked = st.button(
        "✨  Analyze Resume",
        use_container_width=True,
        type="primary"
    )

with clear_col:
    clear_clicked = st.button(
        "Clear",
        use_container_width=True
    )

if clear_clicked:
    st.session_state.pop("analysis_result", None)
    st.rerun()


# --------------------------------------------------
# RUN ANALYSIS
# --------------------------------------------------

if analyze_clicked:
    if uploaded_file is None:
        st.error("Please upload your resume first.")

    elif not job_description.strip():
        st.error("Please enter the job description.")

    else:
        with st.spinner("Reading your resume and analyzing keywords..."):
            resume_text, extraction_error = extract_resume_text(uploaded_file)

            if extraction_error:
                st.error(extraction_error)
                st.session_state.pop("analysis_result", None)

            else:
                result = calculate_analysis(resume_text, job_description)

                if not result["job_skills"]:
                    st.warning(
                        "No known skills were identified in the job description. "
                        "The score may not be meaningful. Try including specific "
                        "technical skills or job requirements."
                    )

                st.session_state["analysis_result"] = result
                st.success("Resume analysis completed!")


# --------------------------------------------------
# RESULTS DASHBOARD
# --------------------------------------------------

result = st.session_state.get("analysis_result")

st.divider()

st.markdown(
    '<div class="section-heading">Analysis Dashboard</div>',
    unsafe_allow_html=True
)

st.markdown(
    '<div class="section-subtitle">'
    'Your resume-to-job-description comparison and improvement insights.'
    '</div>',
    unsafe_allow_html=True
)

if result:
    score = result["score"]
    matched = result["matched"]
    missing = result["missing"]

    # METRIC CARDS

    metric1, metric2, metric3, metric4 = st.columns(4, gap="medium")

    with metric1:
        st.markdown(f"""
        <div class="metric-card">
            <div class="metric-label">ATS-style Match Score</div>
            <div class="metric-value">{score}%</div>
        </div>
        """, unsafe_allow_html=True)

    with metric2:
        st.markdown(f"""
        <div class="metric-card">
            <div class="metric-label">Matched Skills</div>
            <div class="metric-value">{len(matched)}</div>
        </div>
        """, unsafe_allow_html=True)

    with metric3:
        st.markdown(f"""
        <div class="metric-card">
            <div class="metric-label">Missing Skills</div>
            <div class="metric-value">{len(missing)}</div>
        </div>
        """, unsafe_allow_html=True)

    with metric4:
        st.markdown(f"""
        <div class="metric-card">
            <div class="metric-label">Resume Word Count</div>
            <div class="metric-value">{result["resume_words"]}</div>
        </div>
        """, unsafe_allow_html=True)

    # GAUGE AND CHART

    st.markdown("### Match overview")

    gauge_col, chart_col = st.columns([1, 1], gap="large")

    with gauge_col:
        st.plotly_chart(
            create_gauge(score),
            use_container_width=True,
            config={"displayModeBar": False}
        )

        if score >= 75:
            st.success("Strong keyword alignment. Review the remaining gaps.")
        elif score >= 50:
            st.warning("Moderate alignment. Review the missing skills.")
        else:
            st.info("Low keyword alignment. Review the job requirements and resume content.")

    with chart_col:
        st.markdown("#### Skill coverage")
        st.plotly_chart(
            create_skill_chart(matched, missing),
            use_container_width=True,
            config={"displayModeBar": False}
        )

        st.caption(
            f"Job-description word coverage: {result['word_coverage']}%. "
            "This is a simple text comparison, not an ATS ranking."
        )

    # TABS

    st.markdown("### Detailed feedback")

    tab_strengths, tab_fixes, tab_warnings, tab_keywords = st.tabs([
        "✅ Strengths",
        "🛠️ Improvements",
        "⚠️ Warnings",
        "🔎 Keyword gaps"
    ])

    with tab_strengths:
        if result["strengths"]:
            for item in result["strengths"]:
                st.markdown(f"**{item['title']}**")
                st.write(item["description"])
                st.divider()
        else:
            st.info("No strengths identified yet.")

        if matched:
            st.markdown("#### Matched skills")

            skill_columns = st.columns(3)

            for index, skill in enumerate(matched):
                with skill_columns[index % 3]:
                    st.success(f"✓ {skill}")

    with tab_fixes:
        if result["fixes"]:
            for item in result["fixes"]:
                st.markdown(f"**{item['title']}**")
                st.write(item["description"])
                st.divider()
        else:
            st.success("No keyword-related improvement suggestions were generated.")

        st.markdown("#### General recommendations")

        st.write("- Use clear headings for education, skills, projects, and experience.")
        st.write("- Describe projects using actions, tools, and measurable outcomes.")
        st.write("- Tailor relevant resume content to the job description.")
        st.write("- Add a skill only when you genuinely have that skill.")

    with tab_warnings:
        if result["warnings"]:
            for item in result["warnings"]:
                st.warning(f"**{item['title']}**\n\n{item['description']}")
        else:
            st.success("No basic extraction warnings were detected.")

        st.caption(
            "This tool does not verify employment history, factual accuracy, "
            "or compatibility with a specific employer's ATS."
        )

    with tab_keywords:
        keyword_col1, keyword_col2 = st.columns(2, gap="large")

        with keyword_col1:
            st.markdown("#### Matched keywords")

            if matched:
                for skill in matched:
                    st.markdown(f"- 🟢 {skill}")
            else:
                st.info("No matching skills found.")

        with keyword_col2:
            st.markdown("#### Missing keywords")

            if missing:
                for skill in missing:
                    st.markdown(f"- 🔴 {skill}")
            else:
                st.success("No missing skills from the detected skill list.")

        st.markdown("#### All skills found in resume")

        if result["resume_skills"]:
            st.write(", ".join(result["resume_skills"]))
        else:
            st.write("No skills from the current built-in skill list were detected.")

        st.markdown("#### Skills detected in job description")

        if result["job_skills"]:
            st.write(", ".join(result["job_skills"]))
        else:
            st.write("No skills from the current built-in skill list were detected.")

    # RESUME PREVIEW

    st.markdown("### Resume preview")

    with st.expander("View extracted resume text"):
        st.text_area(
            "Extracted content",
            value=result["resume_text"],
            height=350,
            disabled=True,
            label_visibility="collapsed"
        )

    # DOWNLOADABLE REPORT

    st.markdown("### Export your report")

    report_lines = [
        "RESUMEAI ANALYSIS REPORT",
        "=" * 35,
        "",
        f"Resume word count: {result['resume_words']}",
        f"ATS-style keyword match score: {score}%",
        f"Matched skills: {len(matched)}",
        f"Missing skills: {len(missing)}",
        "",
        "MATCHED SKILLS",
        "-" * 20,
        ", ".join(matched) if matched else "None identified",
        "",
        "MISSING SKILLS",
        "-" * 20,
        ", ".join(missing) if missing else "None identified",
        "",
        "SKILLS FOUND IN RESUME",
        "-" * 20,
        ", ".join(result["resume_skills"]) or "None identified",
        "",
        "SKILLS FOUND IN JOB DESCRIPTION",
        "-" * 20,
        ", ".join(result["job_skills"]) or "None identified",
        "",
        "GENERAL RECOMMENDATIONS",
        "-" * 20,
        "Use relevant job-description terms only when they accurately describe your experience.",
        "Use clear resume headings and quantify project outcomes when possible.",
        "",
        "IMPORTANT NOTE",
        "-" * 20,
        "This score is a basic keyword match estimate, not an official ATS score.",
        "Results depend on the built-in skill list and text extraction quality."
    ]

    report_text = "\n".join(report_lines)

    st.download_button(
        label="⬇️ Download Analysis Report",
        data=report_text,
        file_name="resume_analysis_report.txt",
        mime="text/plain",
        use_container_width=True
    )

else:
    empty1, empty2, empty3 = st.columns(3)

    with empty1:
        st.markdown("""
        <div class="metric-card">
            <div class="metric-label">ATS-style Match Score</div>
            <div class="metric-value">--</div>
        </div>
        """, unsafe_allow_html=True)

    with empty2:
        st.markdown("""
        <div class="metric-card">
            <div class="metric-label">Matched Skills</div>
            <div class="metric-value">--</div>
        </div>
        """, unsafe_allow_html=True)

    with empty3:
        st.markdown("""
        <div class="metric-card">
            <div class="metric-label">Missing Skills</div>
            <div class="metric-value">--</div>
        </div>
        """, unsafe_allow_html=True)

    st.info(
        "Upload your resume, paste a job description, and click "
        "**Analyze Resume** to see your results here."
    )


# --------------------------------------------------
# FOOTER
# --------------------------------------------------

st.markdown("""
<div class="footer">
    ResumeAI Analyzer · Built with Streamlit, Python, PyPDF2, Pandas and Plotly
    <br>
    Analysis is based on extracted text and a built-in skill list.
</div>
""", unsafe_allow_html=True)
